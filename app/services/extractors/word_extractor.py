from docx import Document

from app.services.extractors.base_extractor import BaseExtractor


class WordExtractor(BaseExtractor):

    def extract(self, file):
        document = Document(file)

        text_parts = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text.strip())

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                if any(cells):
                    text_parts.append(" | ".join(cells))

        return {
            "text": "\n".join(text_parts),
            "source_type": "word"
        }