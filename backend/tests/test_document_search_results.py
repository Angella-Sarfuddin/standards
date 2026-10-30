import unittest
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from app.services.document_search import result_helpers
from app.services.document_search.enrichment_builder import (
    committee_display_codes,
)
from app.services.document_search_service import DocumentSearchService
from app.services.export_service import ExportService


class _SearchService:
    @staticmethod
    def _not_found_reason(result):
        return result_helpers.not_found_reason(result)


def _resolved_result(source_identity, returned):
    return {
        "requested_sdo": "BIS",
        "requested_standard": returned,
        "matched": True,
        "status": "AVAILABLE",
        "standard": {
            "Standard_id": "BIS-3025-PT24-2021",
            "sdo_id": 1,
            "standardno": returned,
            "display_stdNo": returned,
        },
        "_source_request_identity": source_identity,
    }


def _not_found_result(source_identity, requested):
    return {
        "requested_sdo": "BIS",
        "requested_standard": requested,
        "matched": False,
        "status": "NOT FOUND",
        "standard": None,
        "actions": {"manual_search": True},
        "_source_request_identity": source_identity,
    }


class ResolvedEditionReportingTests(unittest.TestCase):
    def setUp(self):
        self.old_edition = "IS 3025 PT 24 : 1986"
        self.current_edition = "IS 3025 PT 24 : 2021"
        self.source_identity = (
            "customer-list.xlsx",
            8,
            "IS3025PT24:1986",
        )

    def normalize(self, results):
        return result_helpers.deduplicate_results(results, _SearchService)

    def test_resolved_older_edition_is_not_reported_not_found(self):
        resolved = _resolved_result(
            self.source_identity,
            self.current_edition,
        )
        stale_unmatched = _not_found_result(
            self.source_identity,
            self.old_edition,
        )

        results = self.normalize([resolved, stale_unmatched])
        not_found = [result for result in results if not result["matched"]]

        self.assertEqual(len(results), 1)
        self.assertEqual(
            results[0]["standard"]["display_stdNo"],
            self.current_edition,
        )
        self.assertEqual(not_found, [])

    def test_genuinely_missing_standard_remains_not_found(self):
        missing = _not_found_result(
            ("customer-list.xlsx", 9, "IS3025PT99:2015"),
            "IS 3025 PT 99 : 2015",
        )

        results = self.normalize([missing])

        self.assertEqual(len(results), 1)
        self.assertFalse(results[0]["matched"])
        self.assertEqual(results[0]["status"], "NOT FOUND")

    def test_similar_but_unrelated_standard_does_not_resolve_request(self):
        similar_found = _resolved_result(
            ("customer-list.xlsx", 10, "IS3025PT240:2021"),
            "IS 3025 PT 240 : 2021",
        )
        requested_missing = _not_found_result(
            ("customer-list.xlsx", 10, "IS3025PT24:1986"),
            self.old_edition,
        )

        results = self.normalize([similar_found, requested_missing])

        self.assertEqual(len(results), 2)
        self.assertTrue(results[0]["matched"])
        self.assertFalse(results[1]["matched"])
        self.assertEqual(results[1]["requested_standard"], self.old_edition)

    def test_ambiguous_publisher_request_remains_unresolved(self):
        service = DocumentSearchService.__new__(DocumentSearchService)
        service.enrichment_repository = SimpleNamespace(
            search_master_candidates_across_sdos=lambda _core: [
                (
                    SimpleNamespace(
                        sdo_id=218,
                        standardno="IRC SP 75",
                        display_stdNo="IRC SP 75 : 2008",
                    ),
                    "IRC",
                ),
                (
                    SimpleNamespace(
                        sdo_id=52,
                        standardno="MSS SP 75",
                        display_stdNo="MSS SP 75 : 2025",
                    ),
                    "MSS",
                ),
            ]
        )

        self.assertIsNone(service._infer_unique_catalog_sdo("SP", "75"))

        ambiguous = _not_found_result(
            ("customer-list.xlsx", 12, "SP75"),
            "SP 75",
        )
        ambiguous.update(
            requested_sdo="SP",
            status="SDO NOT FOUND",
        )

        results = self.normalize([ambiguous])

        self.assertEqual(len(results), 1)
        self.assertFalse(results[0]["matched"])
        self.assertEqual(results[0]["status"], "SDO NOT FOUND")
        self.assertIsNone(results[0]["standard"])
        self.assertTrue(results[0]["actions"]["manual_search"])

    def test_explicit_amendment_request_remains_unresolved(self):
        base_found = _resolved_result(
            ("customer-list.xlsx", 11, "IS3025PT24:1986"),
            self.current_edition,
        )
        amendment_missing = _not_found_result(
            ("customer-list.xlsx", 11, "IS3025PT24:1986AMD1"),
            "IS 3025 PT 24 : 1986 AMD 1",
        )
        amendment_missing["status_detail"] = (
            "Revision/amendment requests are excluded from standard matching."
        )

        results = self.normalize([base_found, amendment_missing])

        self.assertEqual(len(results), 2)
        self.assertTrue(results[0]["matched"])
        self.assertFalse(results[1]["matched"])
        self.assertIn("amendment", results[1]["status_detail"].casefold())

    def test_successful_request_does_not_hide_another_missing_request(self):
        resolved = _resolved_result(
            self.source_identity,
            self.current_edition,
        )
        missing = _not_found_result(
            ("customer-list.xlsx", 9, "IS3025PT99:2015"),
            "IS 3025 PT 99 : 2015",
        )

        results = self.normalize([resolved, missing])

        self.assertEqual(len(results), 2)
        self.assertTrue(results[0]["matched"])
        self.assertFalse(results[1]["matched"])
        self.assertEqual(results[1]["requested_standard"], "IS 3025 PT 99 : 2015")

    def test_exact_version_match_remains_found(self):
        exact_source_identity = (
            "customer-list.xlsx",
            8,
            "IS3025PT24:2021",
        )
        exact = _resolved_result(
            exact_source_identity,
            self.current_edition,
        )

        results = self.normalize([exact])

        self.assertEqual(len(results), 1)
        self.assertTrue(results[0]["matched"])
        self.assertEqual(
            results[0]["requested_standard"],
            self.current_edition,
        )

    def test_older_edition_on_another_row_is_not_reported_not_found(self):
        resolved = _resolved_result(
            ("customer-list.xlsx", 4, "IS3025PT24:2021"),
            self.current_edition,
        )
        older = _not_found_result(
            ("customer-list.xlsx", 8, "IS3025PT24:1986"),
            self.old_edition,
        )

        results = self.normalize([resolved, older])
        workbook_bytes = ExportService.export_not_found_excel(results)
        workbook = load_workbook(BytesIO(workbook_bytes.getvalue()), read_only=True)
        try:
            standards = [
                row[1]
                for row in workbook["Not Found"].iter_rows(
                    min_row=2,
                    values_only=True,
                )
            ]
        finally:
            workbook.close()

        self.assertEqual(len(results), 1)
        self.assertTrue(results[0]["matched"])
        self.assertEqual(
            results[0]["standard"]["display_stdNo"],
            self.current_edition,
        )
        self.assertEqual(standards, [])

    def test_different_part_remains_not_found(self):
        resolved = _resolved_result(
            ("customer-list.xlsx", 1, "IEC623051:2024"),
            "IEC 62305-1:2024 (EN-FR)",
        )
        resolved["requested_sdo"] = "IEC"
        other_part = _not_found_result(
            ("customer-list.xlsx", 2, "IEC623052:2024"),
            "IEC 62305-2:2024 (EN-FR)",
        )
        other_part["requested_sdo"] = "IEC"

        results = self.normalize([resolved, other_part])

        self.assertEqual(len(results), 2)
        self.assertFalse(results[1]["matched"])
        self.assertEqual(
            results[1]["requested_standard"],
            "IEC 62305-2:2024 (EN-FR)",
        )

    def test_year_newer_than_found_edition_remains_not_found(self):
        resolved = _resolved_result(
            ("customer-list.xlsx", 1, "IEC61000429:2000"),
            "IEC 61000-4-29:2000 (EN-FR)",
        )
        resolved["requested_sdo"] = "IEC"
        newer = _not_found_result(
            ("customer-list.xlsx", 2, "IEC61000429:2026"),
            "IEC 61000-4-29:2026",
        )
        newer["requested_sdo"] = "IEC"

        results = self.normalize([resolved, newer])

        self.assertEqual(len(results), 2)
        self.assertFalse(results[1]["matched"])
        self.assertEqual(
            results[1]["requested_standard"],
            "IEC 61000-4-29:2026",
        )

    def test_filtered_results_keep_not_found_workbook_consistent(self):
        resolved = _resolved_result(
            self.source_identity,
            self.current_edition,
        )
        stale_unmatched = _not_found_result(
            self.source_identity,
            self.old_edition,
        )
        genuinely_missing = _not_found_result(
            ("customer-list.xlsx", 9, "IS3025PT99:2015"),
            "IS 3025 PT 99 : 2015",
        )
        results = self.normalize([resolved, stale_unmatched, genuinely_missing])

        workbook_bytes = ExportService.export_not_found_excel(results)
        workbook = load_workbook(BytesIO(workbook_bytes.getvalue()), read_only=True)
        try:
            standards = [
                row[1]
                for row in workbook["Not Found"].iter_rows(
                    min_row=2,
                    values_only=True,
                )
            ]
        finally:
            workbook.close()

        self.assertEqual(standards, ["IS 3025 PT 99 : 2015"])


class CommitteeCodeDisplayTests(unittest.TestCase):
    def test_publishers_show_full_codes_without_titles(self):
        self.assertEqual(
            committee_display_codes(
                "19",
                "Electrical Accessories & Special Equipment (GMB19)",
                "BSI",
            ),
            ["GMB 19"],
        )
        self.assertEqual(
            committee_display_codes(
                "36",
                "CHD 36 -- WATER QUALITY",
                "BIS",
            ),
            ["CHD 36"],
        )
        self.assertEqual(
            committee_display_codes(
                "16/12",
                "E-TD 16 and MHD 12",
                "BIS",
            ),
            ["E-TD 16", "MHD 12"],
        )
        self.assertEqual(
            committee_display_codes(
                "22/7",
                "Adjustable speed electric drive systems",
                "IEC",
            ),
            ["TC 22/SC 7"],
        )
        self.assertEqual(
            committee_display_codes(
                "1",
                "Volume 01.01 Steel--Piping, Tubing, Fittings",
                "ASTM",
            ),
            ["01.01"],
        )
        self.assertIsNone(
            committee_display_codes("65", "ISO/TC 65", "ISO")
        )


if __name__ == "__main__":
    unittest.main()
