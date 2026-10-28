import os
import re
import subprocess
import tempfile
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.services.extractors.base_extractor import BaseExtractor
from app.core.legacy_office import (
    OfficeConversionError,
    OfficeConversionUnavailable,
    converted_office_file,
)
from app.core.standard_expression_parser import StandardExpressionParser


# A column whose header is one of these holds the standard, one per row.
_NUMBER_HEADER = re.compile(
    r"^(?:"
    r"document\s*no\.?|document\s*number|"
    r"standard(?:\s*(?:no\.?|number))?|"
    r"display\s*std\.?\s*no\.?|designation|"
    r"std\.?\s*no\.?|specification\s*no\.?"
    r")$",
    re.IGNORECASE,
)
_SKIP_CELL = re.compile(
    r"^(?:-|\u2013|n/?a|nil|"
    r"(?:\d+\s*)?(?:st|nd|rd|th)\s+edition|"
    r"edition|effective\s+date|errata|addendum|remarks"
    r")$",
    re.IGNORECASE,
)


class WordExtractor(BaseExtractor):

    def extract(self, file, filename=None):
        try:
            if hasattr(file, "seek"):
                file.seek(0)
            document = Document(file)
        except (PackageNotFoundError, BadZipFile, ValueError):
            if hasattr(file, "seek"):
                file.seek(0)
            legacy = self._extract_legacy_doc(file, filename)
            if isinstance(legacy, list):
                return legacy
            return {
                "text": legacy,
                "source_type": "word",
            }

        return self._extract_docx_document(document)

    @classmethod
    def _extract_docx_document(cls, document):
        tables = []
        for table in document.tables:
            rows = []
            for row in table.rows:
                rows.append([cls._clean_cell(cell.text) for cell in row.cells])
            if rows:
                tables.append(rows)

        selected = cls._standards_from_tables(tables)
        if selected:
            # Keep the Word extractor contract consistent with the legacy
            # conversion path: callers expect canonical request dictionaries,
            # not a bare list of designation strings.  Returning strings here
            # made DocumentService pass each string to extract_from_rows(),
            # which expects mappings and failed with ``'str' object has no
            # attribute 'items'`` for table-based Word documents.
            return cls._as_rows(selected)

        text_parts = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text.strip())
        for rows in tables:
            for row in rows:
                if any(row):
                    text_parts.append(" | ".join(row))
        return {
            "text": "\n".join(text_parts),
            "source_type": "word",
        }

    def _extract_legacy_doc(self, file, filename=None):
        """Convert binary .doc to docx, retaining table-aware extraction."""
        try:
            with converted_office_file(
                file,
                filename or getattr(file, "name", "upload.doc"),
                ".docx",
            ) as converted:
                return self._extract_docx_document(Document(str(converted)))
        except OfficeConversionUnavailable:
            # Older installations may have Microsoft Word but not LibreOffice.
            return self._extract_legacy_doc_with_word(file)
        except (OfficeConversionError, PackageNotFoundError, BadZipFile, ValueError) as conversion_error:
            try:
                return self._extract_legacy_doc_with_word(file)
            except ValueError as word_error:
                raise ValueError(
                    "This older .doc file could not be read. "
                    f"LibreOffice: {conversion_error} Word: {word_error}"
                ) from conversion_error

    def _extract_legacy_doc_with_word(self, file):
        """Read a binary .doc through Word as a fallback."""
        path, cleanup = self._materialize_doc(file)
        escaped = path.replace("'", "''")
        command = (
            "$word = New-Object -ComObject Word.Application; "
            "$word.Visible = $false; "
            "$word.DisplayAlerts = 0; "
            "[Console]::OutputEncoding = [System.Text.Encoding]::UTF8; "
            "function Clean([string]$t) { "
            "if (-not $t) { return '' }; "
            "return (($t -replace '[\\r\\a]', ' ') -replace '\\s+', ' ').Trim() "
            "}; "
            "try { "
            f"$doc = $word.Documents.Open('{escaped}', $false, $true); "
            "$header = '^(document\\s*no\\.?|document\\s*number|standard(\\s*(no\\.?|number))?|display\\s*std\\.?\\s*no\\.?|designation|std\\.?\\s*no\\.?|specification\\s*no\\.?)$'; "
            "if ($doc.Tables.Count -eq 0) { "
            "Write-Output '---TEXT---'; Write-Output $doc.Content.Text "
            "} else { "
            "$found = $false; "
            "foreach ($table in $doc.Tables) { "
            "$col = 0; "
            "$limit = [Math]::Min(3, $table.Rows.Count); "
            "for ($r = 1; $r -le $limit -and $col -eq 0; $r++) { "
            "for ($c = 1; $c -le $table.Columns.Count; $c++) { "
            "try { $t = Clean $table.Cell($r, $c).Range.Text } catch { $t = '' }; "
            "if ($t -match $header) { $col = $c; break } "
            "} }; "
            "if ($col -eq 0) { continue }; "
            "$found = $true; "
            "Write-Output '---NUMBERS---'; "
            "for ($r = 1; $r -le $table.Rows.Count; $r++) { "
            "try { $t = Clean $table.Cell($r, $col).Range.Text } catch { $t = '' }; "
            "Write-Output $t "
            "} }; "
            "if (-not $found) { Write-Output '---TEXT---'; Write-Output $doc.Content.Text } "
            "} "
            "$doc.Close($false) "
            "} finally { $word.Quit() }"
        )
        try:
            completed = subprocess.run(
                ["powershell", "-NoProfile", "-NonInteractive", "-Command", command],
                capture_output=True,
                check=False,
                timeout=45,
            )
        except subprocess.TimeoutExpired as exc:
            raise ValueError(
                "Microsoft Word timed out while reading this .doc file. "
                "Install/configure LibreOffice to use the automatic converter."
            ) from exc
        except OSError as exc:
            raise ValueError(
                "Microsoft Word automation is unavailable for this .doc file. "
                "Install/configure LibreOffice to use the automatic converter."
            ) from exc
        finally:
            if cleanup:
                os.remove(path)

        if completed.returncode != 0:
            detail = (completed.stderr or completed.stdout or b"").decode(
                "utf-8",
                errors="replace",
            ).strip()
            raise ValueError(
                "This Word file is the older .doc format and could not be read. "
                + detail
            )

        raw = completed.stdout.decode("utf-8", errors="replace")
        if "---TEXT---" in raw and "---NUMBERS---" not in raw:
            text = raw.split("---TEXT---", 1)[1]
            text = text.replace("\r\x07", "\n").replace("\x07", "\n").replace("\r", " ")
            return text
        numbers = []
        for line in raw.splitlines():
            if line.strip() == "---NUMBERS---":
                continue
            value = self._clean_cell(line)
            if not value or _SKIP_CELL.match(value) or _NUMBER_HEADER.match(value):
                continue
            numbers.append(value)
        return self._as_rows(numbers)

    @staticmethod
    def _as_rows(numbers):
        parser = StandardExpressionParser()
        rows = []
        for number in numbers:
            rows.extend(parser.designation_cell(number))
        return rows

    @classmethod
    def _standards_from_tables(cls, tables):
        numbers = []
        for rows in tables:
            column = cls._number_column(rows)
            if column is None:
                continue
            for row in rows:
                if column >= len(row):
                    continue
                value = cls._clean_cell(row[column])
                if not value or _SKIP_CELL.match(value) or _NUMBER_HEADER.match(value):
                    continue
                numbers.append(value)
        return numbers

    @staticmethod
    def _number_column(rows):
        for row in rows[:3]:
            for index, cell in enumerate(row):
                if _NUMBER_HEADER.match(cell or ""):
                    return index
        return None

    @staticmethod
    def _clean_cell(value):
        text = str(value or "")
        text = text.replace("\u00a0", " ").replace("\ufffd", "-")
        text = text.replace("\u2010", "-").replace("\u2011", "-")
        text = text.replace("\u2012", "-").replace("\u2013", "-").replace("\u2014", "-")
        text = re.sub(r"\s*/\s*", "/", text)
        text = re.sub(r"\s+", " ", text).strip(" |")
        return text

    @staticmethod
    def _materialize_doc(file):
        if isinstance(file, (str, os.PathLike)) and os.path.isfile(file):
            return os.path.abspath(file), False

        name = getattr(file, "name", None)
        if isinstance(name, str) and os.path.isfile(name):
            return os.path.abspath(name), False

        raw = file.read() if hasattr(file, "read") else bytes(file)
        temporary = tempfile.NamedTemporaryFile(delete=False, suffix=".doc")
        temporary.write(raw)
        temporary.close()
        return temporary.name, True
