
import re
from html import unescape
import zipfile
from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from app.services.grouping_service import GroupingService
from app.utils.api_classification_filter import filter_api_classifications


class ExportService:
    QUOTATION_COLUMNS = [
        "SDO",
        "Standard",
        "Matched Standard",
        "Title",
        "Standard Year",
        "Edition Update",
        "Currency",
        "Price",
        "TC",
    ]
    GROUPING_COLUMNS = [
        "SDO",
        "TC",
        "Type",
        "Standards Count",
        "Individual Total",
        "Group Price",
        "Potential Saving",
    ]
    LATEST_YEAR_COLUMNS = [
        "SDO",
        "Search Standard",
        "Latest Available Year",
        "Matched Catalog Standard",
        "Title",
    ]
    NOT_FOUND_COLUMNS = [
        "SDO",
        "Standard",
        "Status",
        "Superseded By",
    ]

    @staticmethod
    def export_results(response):

        results = response.get("results", [])
        grouping = response.get("grouping", {})

        quotation_rows = []

        for result in results:

            if not result.get("matched"):
                continue

            standard = result.get("standard")

            if not standard:
                continue

            enrichment = standard.get(
                "enrichment",
                {}
            )

            classifications = enrichment.get(
                "classifications",
                []
            )
            classifications = filter_api_classifications(
                standard.get("standardno"),
                classifications,
            )

            price = enrichment.get("price")
            requested_standard = str(
                result.get("requested_standard") or ""
            )
            catalog_designation = " ".join(
                str(value or "")
                for value in (
                    standard.get("display_stdNo"),
                    standard.get("standardno"),
                )
            )
            is_bis_standard = ExportService._is_bis_standard(
                result.get("requested_sdo"),
                standard.get("display_stdNo") or standard.get("standardno"),
            )
            has_international_equivalent = bool(
                re.search(
                    r"(?i)\b(?:ISO|IEC|IEEE)\b",
                    f"{requested_standard} {catalog_designation}",
                )
            )
            is_simple_is_edition = bool(
                is_bis_standard
                and re.match(
                    r"(?i)\s*(?:IS|BIS)\b",
                    f"{requested_standard} {catalog_designation}",
                )
                and not has_international_equivalent
            )
            # Plain IS editions are free; discard any paid alias amount.
            # ISO/IEC/IEEE adoptions retain their listed database price.
            if is_simple_is_edition:
                price = 0
            standard_year = standard.get("standardyear")
            if standard_year is None:
                standard_year = enrichment.get("year")

            currency = str(
                standard.get("currency")
                or enrichment.get("currency")
                or ""
            ).strip() or None

            quotation_row = {
                "SDO": ExportService._export_sdo(
                    result.get("requested_sdo"),
                    standard.get("display_stdNo") or standard.get("standardno"),
                ),
                # Display the resolved current catalog designation. This keeps
                # request-only COR/AMD suffixes out of the workbook.
                "Standard": (
                    standard.get("display_stdNo")
                    or standard.get("standardno")
                    or ExportService._strip_revision_suffix(
                        result.get("requested_standard")
                    )
                ),
                "Matched Standard": (
                    standard.get("display_stdNo")
                    or standard.get("standardno")
                ),
                "Title": standard.get("title"),
                "Standard Year": standard_year,
                "Edition Update": result.get("status_detail")
                if result.get("status_detail")
                else None,
                "Currency": currency,
                "Price": price,
            }

            # Keep one quotation row per found request. A standard may belong
            # to several committees; list their catalogue descriptions in the
            # last column instead of multiplying the standard across rows.
            tc_descriptions = []
            requested_sdo = re.sub(
                r"[^A-Z0-9]",
                "",
                str(result.get("requested_sdo") or "").upper(),
            )
            if requested_sdo.startswith("ASME"):
                # ASME designations/volumes are not technical committees.
                classifications = []
            seen_classifications = set()
            for classification in classifications:
                code = str(classification.get("code") or "").strip()
                description = str(
                    classification.get("description") or ""
                ).strip()
                code = ExportService._clean_classification_text(code)
                description = ExportService._clean_classification_text(
                    description
                )
                key = (code.casefold(), description.casefold())
                if not code and not description or key in seen_classifications:
                    continue
                seen_classifications.add(key)
                tc_text = description or code
                if tc_text and tc_text not in tc_descriptions:
                    tc_descriptions.append(tc_text)

            quotation_rows.append(
                {
                    **quotation_row,
                    "TC": ", ".join(tc_descriptions) or None,
                }
            )

        quotation_df = pd.DataFrame(
            quotation_rows,
            columns=ExportService.QUOTATION_COLUMNS,
        )
        quotation_df = ExportService._sort_by_sdo(quotation_df)

        grouping_rows = []

        if grouping.get("enabled"):

            for group in grouping.get("groups", []):

                grouping_rows.append({
                    "SDO": ExportService._export_sdo(group.get("sdo")),
                    "TC": group.get("tc"),
                    "Type": group.get(
                        "type"
                    ),
                    "Standards Count": group.get(
                        "count"
                    ),
                    "Individual Total": group.get(
                        "individual_total"
                    ),
                    "Group Price": group.get(
                        "group_price"
                    ),
                    "Potential Saving": group.get(
                        "potential_saving"
                    )
                })

        grouping_df = pd.DataFrame(
            grouping_rows,
            columns=ExportService.GROUPING_COLUMNS,
        )
        grouping_df = ExportService._sort_by_sdo(grouping_df)

        output = BytesIO()

        with pd.ExcelWriter(
            output,
            engine="openpyxl"
        ) as writer:

            quotation_df.to_excel(
                writer,
                index=False,
                sheet_name="Quotation"
            )

            grouping_df.to_excel(
                writer,
                index=False,
                sheet_name="Grouping Summary"
            )

        output.seek(0)
        ExportService._format_workbook(output)
        output.seek(0)
        return output

    @staticmethod
    def export_latest_years(results):
        """Export requested standards with their latest matched catalog year."""
        request_rows = {}
        request_order = []

        for result in results or []:
            sdo = str(result.get("requested_sdo") or "").strip()
            requested = str(result.get("requested_standard") or "").strip()
            if not sdo and not requested:
                continue

            key = (sdo.casefold(), requested.casefold())
            if key not in request_rows:
                request_order.append(key)
                request_rows[key] = {
                    "SDO": ExportService._export_sdo(sdo, requested),
                    "Search Standard": ExportService._strip_revision_suffix(
                        requested
                    ),
                    "Latest Available Year": None,
                    "Matched Catalog Standard": None,
                    "Title": None,
                    "_year_rank": -1,
                }

            row = request_rows[key]
            standard = (
                (result.get("standard") or {})
                if result.get("matched")
                else {}
            )
            year = ExportService._available_year(standard)
            if year is not None and year > row["_year_rank"]:
                row["Search Standard"] = (
                    standard.get("display_stdNo")
                    or standard.get("standardno")
                    or row["Search Standard"]
                )
                row["Latest Available Year"] = year
                row["Matched Catalog Standard"] = (
                    standard.get("display_stdNo")
                    or standard.get("standardno")
                )
                row["Title"] = standard.get("title")
                row["_year_rank"] = year
        rows = [
            {
                column: request_rows[key].get(column)
                for column in ExportService.LATEST_YEAR_COLUMNS
            }
            for key in request_order
        ]
        dataframe = pd.DataFrame(
            rows,
            columns=ExportService.LATEST_YEAR_COLUMNS,
        )
        dataframe = ExportService._sort_by_sdo(dataframe)
        output = BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            dataframe.to_excel(
                writer,
                index=False,
                sheet_name="Latest Available Years",
            )
        output.seek(0)
        ExportService._format_workbook(output)
        output.seek(0)
        return output

    @staticmethod
    def _clean_classification_text(value):
        """Decode catalogue entities and normalize whitespace in TC labels."""
        text = unescape(str(value or ""))
        # Some scraped BIS committee descriptions retain Markdown emphasis.
        text = re.sub(r"\*{1,2}", "", text)
        text = re.sub(r"\s+", " ", text).strip()
        return text

    @staticmethod
    def _sort_by_sdo(dataframe):
        """Return a stable ascending SDO sort for exported worksheets."""
        if dataframe is None or dataframe.empty or "SDO" not in dataframe:
            return dataframe
        return dataframe.sort_values(
            by="SDO",
            key=lambda values: values.fillna("").astype(str).str.casefold(),
            kind="stable",
            na_position="last",
        ).reset_index(drop=True)

    @staticmethod
    def _is_bis_standard(sdo, standard_number=None):
        """Identify BIS standards even when the request uses the IS prefix."""
        normalized_sdo = re.sub(r"[^A-Z0-9]", "", str(sdo or "").upper())
        normalized_number = str(standard_number or "").strip().upper()
        return normalized_sdo in {"IS", "BIS"} or bool(
            re.match(r"^(?:IS|BIS)(?:\s|\d|/)", normalized_number)
        )

    @staticmethod
    def _export_sdo(sdo, standard_number=None):
        """Render catalogue aliases using publisher names in SDO columns."""
        text = str(sdo or "").strip()
        normalized_sdo = re.sub(r"[^A-Z0-9]", "", text.upper())
        normalized_number = str(standard_number or "").strip().upper()
        if normalized_sdo == "IS" or (
            normalized_sdo in {"", "-"}
            and ExportService._is_bis_standard(None, standard_number)
        ):
            return "BIS"
        if normalized_sdo in {"BS", "BSI", "PD", "EN"} or re.match(
            r"(?i)^(?:BSI|BS|PD|EN)(?:\s|/|:|-|\d)",
            normalized_number,
        ):
            return "BSI"
        return text

    @staticmethod
    def _strip_revision_suffix(value):
        """Remove trailing corrigendum or amendment notation from a number."""
        text = str(value or "").strip()
        return re.sub(
            r"(?i)(?:/|\+|\s)\s*(?:CORR(?:IGENDUM)?\.?|COR\.?|AMD\.?|AMEND(?:MENT)?\.?)"
            r"\s*(?:NO\.?\s*)?\d+[A-Z]?\b.*$",
            "",
            text,
        ).rstrip(" :/+")

    @staticmethod
    def _classification_label(code, description, owner_sdo=None):
        """Return the full committee designation, such as TC 22/SC 22G."""
        owner = str(owner_sdo or "").strip().upper()
        code = str(code or "").strip()
        # ASTM classification records describe Book of Standards volumes,
        # rather than technical committees. Emit only their volume numbers.
        if owner == "ASTM":
            volumes = re.findall(
                r"(?i)\b(?:BOOK\s+OF\s+STANDARDS\s+)?(?:VOLUME|VOL\.?\s*)\s*(\d{2}\.\d{2})\b",
                description,
            )
            return ", ".join(dict.fromkeys(volumes)) or None
        # IEEE's code is a numeric catalogue category, while its description
        # contains the committee/society name used in the TC column.
        if owner == "IEEE":
            committee_name = re.sub(r"\s+", " ", str(description or "")).strip()
            committee_name = committee_name.split("(", 1)[0].strip()
            return committee_name or None
        # NEMA stores these as named business sectors, not numbered TCs.
        # Its description already includes the short code (e.g. "BS -");
        # repeat that code parenthetically to match the NEMA catalogue label.
        if owner == "NEMA":
            nema_scope = re.match(
                r"^\s*([A-Z]{2})\s*-\s*(.+?)\s*$",
                description,
                re.I,
            )
            if nema_scope:
                code, name = nema_scope.groups()
                code = code.upper()
                return f"{code} - {name} ({code})"
            return description or code or None
        # ISO committee assignments are hierarchical. Match the complete
        # parent/subcommittee path before the generic family matcher below,
        # which would otherwise stop at the first "TC 176" token.
        if owner == "ISO":
            for candidate in (description, code):
                iso_path = re.search(
                    r"\bISO\s*/\s*TC\s*\d+[A-Z]?"
                    r"(?:\s*/\s*(?:SC|WG|JWG|AHG|PT)\s*\d+[A-Z]?)*\.?",
                    candidate,
                    re.I,
                )
                if iso_path:
                    return re.sub(r"\s+", " ", iso_path.group(0)).strip().rstrip(".")

                # Some records omit the leading ISO/ prefix but still store
                # the full TC/SC assignment, e.g. "TC 176/SC 2".
                iso_tc_path = re.search(
                    r"(?<![A-Za-z0-9])TC\s*[- ]?\s*\d+[A-Z]?"
                    r"(?:\s*/\s*(?:SC|WG|JWG|AHG|PT)\s*[- ]?\s*\d+[A-Z]?)+",
                    candidate,
                    re.I,
                )
                if iso_tc_path:
                    normalized_path = re.sub(
                        r"(?i)\bTC\s*[- ]?\s*", "TC ", iso_tc_path.group(0)
                    )
                    normalized_path = re.sub(
                        r"(?i)\b(SC|WG|JWG|AHG|PT)\s*[- ]?\s*",
                        r"\1 ",
                        normalized_path,
                    )
                    return re.sub(r"\s+", " ", normalized_path).strip()
        # Prefer a complete designation in the description over a shorter
        # taxonomy code, while retaining the generic code fallbacks below.
        iso_path = re.search(
            r"\bISO\s*/\s*TC\s*\d+[A-Z]?"
            r"(?:\s*/\s*(?:SC|WG|JWG|AHG|PT)\s*\d+[A-Z]?)*\.?",
            description,
            re.I,
        )
        if iso_path:
            return re.sub(r"\s+", " ", iso_path.group(0)).strip().rstrip(".,")

        full_tc_path = re.search(
            r"(?<![A-Za-z0-9])TC\s*[- ]?\s*\d+[A-Z]?"
            r"(?:\s*[./,:]\s*(?:SC|WG|JWG|AHG|PT)\s*[- ]?\s*\d+[A-Z]?)+",
            description,
            re.I,
        )
        if full_tc_path:
            normalized_path = re.sub(
                r"(?i)\bTC\s*[- ]?\s*", "TC ", full_tc_path.group(0)
            )
            normalized_path = re.sub(
                r"(?i)\b(SC|WG|JWG|AHG|PT)\s*[- ]?\s*",
                r"\1 ",
                normalized_path,
            )
            return re.sub(r"\s+", " ", normalized_path).strip().rstrip(".,")
        # Catalogue classifications sometimes store only the numeric ID in
        # `code` while the description names the actual committee family
        # (for example, SSPC's "Surface Preparation (SC 05)"). Honor an
        # explicit committee designation before interpreting a bare numeric
        # code as TC N. This applies across SDOs, not only NACE.
        explicit_family = re.search(
            r"(?i)(?<![A-Z0-9])"
            r"(TC|SC|TG|WG|JWG|AHG|PT)\s*[- ]?\s*(\d+[A-Z]?)"
            r"(?![A-Z0-9])",
            f"{description} {code}",
        )
        if explicit_family:
            family, suffix = explicit_family.groups()
            return f"{family.upper()} {suffix.upper()}"
        # NACE commonly omits the space and stores a two-digit SC number.
        if owner in {"AMPP-NACE", "ANSI-NACE", "NACE"}:
            nace_sc = re.search(
                r"(?i)(?<![A-Z0-9])SC\s*0*(\d+[A-Z]?)(?![A-Z0-9])",
                f"{description} {code}",
            )
            if nace_sc:
                suffix = nace_sc.group(1)
                if suffix.isdigit():
                    suffix = suffix.zfill(2)
                return f"SC {suffix.upper()}"
        # BSI committee codes use GMB/GBM identifiers. Their linked
        # classification table often stores only the numeric suffix (for
        # example code "78"), so handle BSI before the generic numeric-TC
        # formatter below.
        if owner == "BSI":
            for candidate in (description, code):
                bsi_code = re.search(
                    r"(?<![A-Z0-9])(?:GMB|GBM)\s*\d+[A-Z]?(?![A-Z0-9])",
                    candidate,
                    re.I,
                )
                if bsi_code:
                    return re.sub(r"\s+", "", bsi_code.group(0)).upper()
            if re.fullmatch(r"\d+[A-Z]?", code, re.I):
                return f"GMB{code.upper()}"
            # Don't label a BSI taxonomy value or arbitrary numeric category
            # as a technical committee when its committee identity is absent.
            return None
        # ASME BPVC numeric classifications identify code volumes/sections,
        # not technical committees. Their labels commonly look like BPVC IIAVI.
        if owner.startswith("ASME"):
            return None
        # BIS descriptions commonly append a title after a dash, such as
        # "PGD 1 -- BASIC STANDARDS". Prefer that designation to a numeric
        # catalogue category.
        bis_code = re.match(r"^([A-Za-z]+)\s*[- ]\s*(\d+)\b", description)
        if bis_code and "--" in description:
            return f"{bis_code.group(1).upper()} {bis_code.group(2)}"

        described_committee_codes = re.search(
            r"(?<![A-Za-z0-9])(?:[A-Z]{1,4}(?:-[A-Z]{1,4})?)"
            r"\s*[- ]\s*\d+[A-Z]?"
            r"(?:\s+and\s+(?:[A-Z]{1,4}(?:-[A-Z]{1,4})?)"
            r"\s*[- ]\s*\d+[A-Z]?)*",
            description,
            re.I,
        )
        if described_committee_codes:
            return re.sub(
                r"\s+", " ", described_committee_codes.group(0)
            ).strip()

        normalized_description = re.sub(
            r"[^a-z0-9]+",
            " ",
            description.casefold(),
        ).strip()
        committee_by_scope = {
            "adjustable speed electric drive systems incorporating semiconductor power converters":
                "TC 22.SC 22G",
            "adjustable speed electric power drive systems pds":
                "TC 22.SC 22G",
            "industrial networks": "TC 65.SC 65C",
            "uninterruptible power systems ups": "TC 22.SC 22H",
            "it security techniques": "ISO/IEC JTC 1/SC 27",
            "low voltage switchgear and controlgear": "TC 121.SC 121A",
            "low voltage switchgear and controlgear assemblies":
                "TC 121.SC 121B",
        }
        if normalized_description in committee_by_scope:
            return committee_by_scope[normalized_description]
        # If no specific designation was recognized, a numeric IEC
        # classification code is the TC itself (for example, 99 means TC 99).
        # Keep this fallback before broad free-text scans, where ordinary
        # words such as "and" followed by a voltage can resemble a code.
        if re.fullmatch(r"\d+", code):
            if owner == "ISO/IEC":
                return None
            return f"TC {code}"
        # DIN classification values sometimes arrive in the TC export column
        # as ICS numbers. Preserve the number as supplied and never add a TC
        # prefix to it.
        for candidate in (code, description):
            ics_codes = re.findall(
                r"(?<!\d)\d{2}\.\d{3}(?:\.\d{2})?(?!\d)",
                candidate,
            )
            if ics_codes:
                return ", ".join(dict.fromkeys(ics_codes))
        # CSA's catalogue classifications are subject categories (for example
        # "Energy" and "Petroleum and Natural Gas Industry Systems"), not
        # technical committees. Keep a value only when CSA explicitly labels
        # it as a TC.
        if owner == "CSA" and not re.search(
            r"(?i)\bTC\b", f"{description} {code}"
        ):
            return None
        # IEC's top-level hierarchy number is the responsible TC. Child
        # numbers are taxonomy categories, not SC suffixes, so use a verified
        # scope mapping above when available and otherwise retain the known TC.
        # ISO/IEC numeric paths remain suppressed because they are taxonomy
        # codes and may not identify a JTC or subcommittee.
        iec_hierarchy = re.fullmatch(
            r"(\d+)(?:/\d+)+",
            re.sub(r"\s+", "", code),
        )
        if owner == "IEC" and iec_hierarchy:
            return f"TC {iec_hierarchy.group(1)}"
        if owner == "ISO/IEC" and iec_hierarchy:
            return None

        for candidate in (description, code):
            normalized_code = re.sub(r"\s+", "", str(code or ""))
            iso_hierarchy = re.fullmatch(
                r"(?i)(\d+[A-Z]?)/(\d+[A-Z]?(?:/\d+[A-Z]?)*)",
                normalized_code,
            )
            if iso_hierarchy and (not owner or owner == "ISO"):
                parent_code, descendants = iso_hierarchy.groups()
                child_codes = descendants.split("/")
                labels = [f"ISO/TC {parent_code}"]
                for child_code in child_codes:
                    labels.append(f"SC {child_code}")
                return "/".join(labels)

            iso_path = re.search(
                r"\bISO\s*/\s*TC\s*\d+[A-Z]?"
                r"(?:\s*/\s*(?:SC|WG|JWG|AHG|PT)\s*\d+[A-Z]?)*\.?",
                candidate,
                re.I,
            )
            if iso_path:
                return re.sub(r"\s+", " ", iso_path.group(0)).strip().rstrip(".")

            # Preserve full committee paths, including subcommittees and
            # working groups, wherever they occur in the classification.
            tc_designation = re.search(
                r"(?<![A-Za-z0-9])TC\s*[- ]?\s*\d+[A-Z]?"
                r"(?:\s*[./,:]\s*(?:SC|WG|JWG|AHG|PT)\s*[- ]?\s*\d+[A-Z]?)*",
                candidate,
                re.I,
            )
            if tc_designation:
                return re.sub(
                    r"(?i)\bTC\s*[- ]?\s*",
                    "TC ",
                    tc_designation.group(0),
                ).strip().rstrip(".,")

            # Preserve other committee families and multiple designations,
            # such as BIS "E-TD 16 and MHD 12".
            committee_codes = re.search(
                r"(?<![A-Za-z0-9])(?:[A-Z]{1,4}(?:-[A-Z]{1,4})?)"
                r"\s*[- ]\s*\d+[A-Z]?"
                r"(?:\s+and\s+(?:[A-Z]{1,4}(?:-[A-Z]{1,4})?)"
                r"\s*[- ]\s*\d+[A-Z]?)*",
                candidate,
                re.I,
            )
            if committee_codes:
                return re.sub(r"\s+", " ", committee_codes.group(0)).strip()

        # Normalize stored TC codes while retaining any suffix/subcommittee.
        tc_code = re.fullmatch(r"(?i)TC\s*[- ]?\s*(\d+[A-Z]?)", code)
        if tc_code:
            return f"TC {tc_code.group(1)}"

        # If no recognizable designation is present, use the source code and
        # never append the free-text description.
        return code or description

    @staticmethod
    def export_not_found_excel(results):
        """Build a formatted workbook containing every unmatched request."""
        rows = []
        for result in results or []:
            if result.get("matched"):
                continue

            def clean_field(value, fallback="-"):
                text = str(value or "").strip()
                text = " ".join(text.split())
                return text or fallback

            requested_sdo = clean_field(result.get("requested_sdo"), "")
            requested_standard = clean_field(
                result.get("requested_standard"),
            )
            complete_standard = requested_standard
            if requested_standard:
                has_sdo_prefix = bool(
                    re.match(
                        r"(?i)^(?:ISO(?:/(?:IEC|TS|TR))?|IEC(?:/(?:TR|TS))?|"
                        r"API|ASTM|ASME|NACE|AMPP|AWWA|ASHRAE|IEEE|ACI|JIS|UL|"
                        r"ANSI|BSI|BS|BIS|IS|EN|DIN|CEN|CENELEC|CSA|FM|IGEM|"
                        r"AWS|AISC|ICC|IAPMO|SII|GB|KS|SANS|NEMA|ISA|EI|DNV|"
                        r"PD|SSPC|HI|CIGRE|EIA|ETS|NFPA|SAE|TEMA)\b",
                        requested_standard,
                    )
                )
                if not has_sdo_prefix and requested_sdo:
                    complete_standard = f"{requested_sdo} {requested_standard}"
                complete_standard = ExportService._strip_revision_suffix(
                    complete_standard
                )

            reason = (
                result.get("not_found_reason")
                or ExportService._not_found_reason_fallback(result)
            )
            normalized_reason = str(reason or "").strip().casefold()
            if normalized_reason == "withdrawn":
                display_status = "WITHDRAWN"
            elif normalized_reason == "historical":
                display_status = "HISTORICAL"
            else:
                display_status = "NOT FOUND"

            rows.append(
                {
                    "SDO": ExportService._export_sdo(
                        requested_sdo,
                        requested_standard,
                    ),
                    "Standard": clean_field(complete_standard),
                    "Status": display_status,
                    "Superseded By": (
                        clean_field(
                            result.get("superseded_by")
                            or result.get("replacement_standard")
                        )
                        if result.get("superseded_by")
                        or result.get("replacement_standard")
                        else None
                    ),
                }
            )

        dataframe = pd.DataFrame(
            rows,
            columns=ExportService.NOT_FOUND_COLUMNS,
        )
        dataframe = ExportService._sort_by_sdo(dataframe)
        output = BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            dataframe.to_excel(
                writer,
                index=False,
                sheet_name="Not Found",
            )
        output.seek(0)
        ExportService._format_workbook(output)
        output.seek(0)
        return output

    @staticmethod
    def _not_found_reason_fallback(result):
        """Provide useful reason labels for unmatched export payloads."""
        evidence = " ".join(
            str(value or "")
            for value in (
                result.get("status"),
                result.get("status_detail"),
                result.get("superseded_by"),
                result.get("replacement_standard"),
            )
        ).casefold()
        if re.search(r"\bwithdrawn\b", evidence):
            return "Withdrawn"
        if re.search(
            r"\b(?:historical|historic|superseded|obsolete|archived)\b",
            evidence,
        ):
            return "Historical"
        if re.search(
            r"\b(?:inactive|discontinued|not available)\b",
            evidence,
        ):
            return "Unavailable"
        if "sdo not found" in evidence:
            return "SDO not found"
        return "Not found in catalog"

    @staticmethod
    def export_found_and_not_found_zip(quotation_results, all_results):
        """Bundle found and unmatched standards as Excel workbooks."""
        grouping = GroupingService.group_results(quotation_results or [])
        workbook = ExportService.export_results(
            {
                "results": quotation_results or [],
                "grouping": grouping,
            }
        )
        not_found_workbook = ExportService.export_not_found_excel(
            all_results or []
        )

        output = BytesIO()
        with zipfile.ZipFile(
            output,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as archive:
            archive.writestr(
                "standards_quotation.xlsx",
                workbook.getvalue(),
            )
            archive.writestr(
                "not_found_standards.xlsx",
                not_found_workbook.getvalue(),
            )
        output.seek(0)
        return output

    @staticmethod
    def export_both(response, latest_year_results=None):
        """Bundle the unchanged quotation export and the latest-year export."""
        output = BytesIO()
        quotation = ExportService.export_results(response)
        latest_years = ExportService.export_latest_years(
            (
                latest_year_results
                if latest_year_results is not None
                else response.get("results", [])
            )
        )

        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(
                "standards_quotation.xlsx",
                quotation.getvalue(),
            )
            archive.writestr(
                "standards_latest_years.xlsx",
                latest_years.getvalue(),
            )

        output.seek(0)
        return output

    @staticmethod
    def _available_year(standard):
        enrichment = standard.get("enrichment") or {}
        for value in (
            standard.get("standardyear"),
            enrichment.get("year"),
            standard.get("display_stdNo"),
            standard.get("standardno"),
        ):
            if value is None:
                continue
            try:
                numeric_year = int(float(value))
            except (TypeError, ValueError, OverflowError):
                matches = re.findall(
                    r"(?<!\d)(?:19|20)\d{2}(?!\d)",
                    str(value),
                )
                if matches:
                    return max(int(match) for match in matches)
                continue
            if 1900 <= numeric_year <= 2099:
                return numeric_year

        return None

    @staticmethod
    def _format_workbook(output):
        output.seek(0)
        from openpyxl import load_workbook

        workbook = load_workbook(output)
        header_fill = PatternFill("solid", fgColor="253858")
        header_font = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
        body_font = Font(name="Aptos", size=10, color="253858")
        alternate_fill = PatternFill("solid", fgColor="F3F6FA")
        available_fill = PatternFill("solid", fgColor="E4F4EA")
        warning_fill = PatternFill("solid", fgColor="FFF1D6")
        bottom_border = Border(
            bottom=Side(style="thin", color="DFE5ED")
        )
        wrap_headers = {
            "title",
            "tc description",
                    "tc",
                    "standard",
                    "reason",
                    "reason details",
                    "search standard",
            "matched standard",
            "matched catalog standard",
        }

        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.sheet_view.showGridLines = False
            worksheet.auto_filter.ref = worksheet.dimensions
            worksheet.row_dimensions[1].height = 30

            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(
                    vertical="center",
                    horizontal="left",
                    wrap_text=True,
                )

            for row_index in range(2, worksheet.max_row + 1):
                max_lines = 1
                for cell in worksheet[row_index]:
                    header = str(
                        worksheet.cell(1, cell.column).value or ""
                    ).casefold()
                    cell.font = body_font
                    cell.alignment = Alignment(
                        vertical="top",
                        wrap_text=header in wrap_headers,
                    )
                    cell.border = bottom_border
                    if row_index % 2 == 0:
                        cell.fill = alternate_fill
                    if cell.value is not None:
                        value = str(cell.value)
                        lines = value.count("\n") + 1
                        if header in wrap_headers:
                            available_width = {
                                "title": 52,
                                "tc description": 44,
                                "tc": 54,
                                "standard": 30,
                                "reason": 24,
                                "reason details": 48,
                                "search standard": 32,
                                "matched standard": 32,
                                "matched catalog standard": 32,
                            }.get(header, 32)
                            lines = max(
                                lines,
                                (len(value) + available_width - 1)
                                // available_width,
                            )
                        max_lines = max(max_lines, lines)

                    if header in {"status", "search status"}:
                        status = str(cell.value or "").upper()
                        if status == "AVAILABLE":
                            cell.fill = available_fill
                        elif status in {
                            "NOT FOUND",
                            "HISTORICAL",
                            "WITHDRAWN",
                            "SDO NOT FOUND",
                            "EDITION MISMATCH",
                        }:
                            cell.fill = warning_fill

                worksheet.row_dimensions[row_index].height = min(
                    90,
                    max(21, 18 * max_lines),
                )

            for column_index in range(1, worksheet.max_column + 1):
                header = str(
                    worksheet.cell(1, column_index).value or ""
                ).casefold()
                values = [
                    str(worksheet.cell(row, column_index).value or "")
                    for row in range(1, worksheet.max_row + 1)
                ]
                content_width = max(
                    (len(line) for value in values for line in value.splitlines()),
                    default=0,
                )
                preferred_widths = {
                    "title": 52,
                    "tc description": 44,
                    "tc": 54,
                    "reason": 24,
                    "reason details": 48,
                    "search standard": 32,
                    "matched catalog standard": 32,
                    "standard": 30,
                    "matched standard": 32,
                    "latest available year": 20,
                    "status": 20,
                    "search status": 20,
                }
                width = preferred_widths.get(
                    header,
                    max(12, min(content_width + 3, 28)),
                )
                worksheet.column_dimensions[
                    get_column_letter(column_index)
                ].width = width

        output.seek(0)
        output.truncate(0)
        workbook.save(output)
        output.seek(0)
